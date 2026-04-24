from __future__ import annotations

import json
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, List, Sequence

import pandas as pd
import tensorflow as tf

from database.data_manager import FlappingWingAerodynamics
from model.prssm import PRSSM
from outputs.optimize_hyperparameters_from_comparison import (
    DEFAULT_EXCEL_PATH,
    FEATURE_COLUMNS,
    OUTPUT_COLUMNS,
    build_column_names,
    optimize_hyperparameters,
)
from outputs.outputs import Outputs
from outputs.postprocess_model_outputs import parse_training_log, run_for_output_dir
from training.trainer import Trainer

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None


DEFAULT_AUTOMATION_SHEET = "Automation Runs"
CURRENT_RUN_STATUS_FILE = "current_automation_run.json"
STOP_REQUEST_FILE = "stop_automation_after_current_run.txt"
HEARTBEAT_INTERVAL_SECONDS = 60
DEFAULT_TUNING_RANDOM_SEED = 123
MIN_HISTORY_FOR_SURROGATE = 5
TRAINING_TEST_NUMBER_INTERNAL_COLUMN = "training_test_number"
SUMMARY_ROW_KEY_BY_INTERNAL_HEADER = {
    "training_test_number": "training_test_number",
    "number_of_layers": "number_of_layers",
    "layer_1_activation_function": "layer1_activation_function",
    "layer_1_filters": "layer1_filters",
    "layer_1_kernel_size": "layer1_kernel_size",
    "layer_1_pool_size": "layer1_pool_size",
    "layer_1_stride": "layer1_stride",
    "layer_2_activation_function": "layer2_activation_function",
    "layer_2_filters": "layer2_filters",
    "layer_2_kernel_size": "layer2_kernel_size",
    "layer_2_pool_size": "layer2_pool_size",
    "layer_2_stride": "layer2_stride",
    "tunable_parameters_shuffle": "shuffle",
    "tunable_parameters_lik_seq_length_factor": "lik_seq_length_factor",
    "tunable_parameters_dim_x": "dim_x",
    "tunable_parameters_ind_pnt_num": "ind_pnt_num",
    "tunable_parameters_samples": "samples",
    "tunable_parameters_use_learning_rate_decay": "use_learning_rate_decay",
    "tunable_parameters_learning_rate_decay_steps": "learning_rate_decay_steps",
    "tunable_parameters_learning_rate_decay_rate": "learning_rate_decay_rate",
    "tunable_parameters_learning_rate_decay_staircase": "learning_rate_decay_staircase",
    "tunable_parameters_recog_len": "recog_len",
    "tunable_parameters_recog_model": "recog_model",
    "tunable_parameters_learning_rate": "learning_rate",
    "tunable_parameters_batch_size": "batch_size",
    "tunable_parameters_patience": "patience",
    "tunable_parameters_min_delta": "min_delta",
    "tunable_parameters_early_stopping_split": "early_stopping_split",
    "tunable_parameters_validation_fraction": "validation_fraction",
    "tunable_parameters_zeta_pos": "zeta_pos",
    "tunable_parameters_zeta_mean": "zeta_mean",
    "tunable_parameters_zeta_var": "zeta_var",
    "tunable_parameters_var_x": "var_x",
    "tunable_parameters_var_y": "var_y",
    "tunable_parameters_gp_var": "gp_var",
    "tunable_parameters_gp_len": "gp_len",
    "tunable_parameters_epochs": "epochs",
    "tunable_parameters_epoch": "epoch",
    "total_training_time_seconds": "total_training_time_seconds",
    "total_training_time_hours": "total_training_time_hours",
    "cf_x_test_rmse": "CF_X_RMSE",
    "cf_x_test_r2": "CF_X_R2",
    "cf_y_test_rmse": "CF_Y_RMSE",
    "cf_y_test_r2": "CF_Y_R2",
    "cm_x_test_rmse": "CM_X_RMSE",
    "cm_x_test_r2": "CM_X_R2",
    "cm_y_test_rmse": "CM_Y_RMSE",
    "cm_y_test_r2": "CM_Y_R2",
    "cm_z_test_rmse": "CM_Z_RMSE",
    "cm_z_test_r2": "CM_Z_R2",
    "average_test_rmse": "average_test_rmse",
    "average_test_r2": "average_test_r2",
    "average_test_training_classification": "training_classification",
    "notes": "notes",
}


def get_repo_root() -> Path:
    return Path(__file__).resolve().parents[1]


def get_data_root() -> Path:
    return get_repo_root().parent / "Data"


def get_default_input_dir() -> Path:
    return get_data_root() / "Input"


def get_default_output_root() -> Path:
    return get_data_root() / "Output"


def get_automation_status_root() -> Path:
    status_root = get_default_output_root()
    status_root.mkdir(parents=True, exist_ok=True)
    return status_root


def get_default_experiment_config() -> Dict[str, Any]:
    return {
        "seq_len": 470,
        "lik_seq_length_factor": 0.8,
        "seq_stride": 470,
        "num_gpus": 1,
        "epochs": 1000,
        "test_data": True,
        "early_stopping": True,
        "early_stopping_split": "validation",
        "validation_fraction": 0.1,
        "batch_size": 14,
        "shuffle": 10000,
        "dim_x": 12,
        "ind_pnt_num": 100,
        "samples": 30,
        "learning_rate": 0.03,
        "use_learning_rate_decay": False,
        "learning_rate_decay_steps": 1000,
        "learning_rate_decay_rate": 0.96,
        "learning_rate_decay_staircase": True,
        "patience": 2,
        "min_delta": 3.162278,
        "recog_len": 60,
        "recog_model": "zeros",
        "number_of_layers": 1,
        "layer1_activation_function": "sigmoid",
        "layer1_filters": 6,
        "layer1_kernel_size": 3,
        "layer1_pool_size": 2,
        "layer1_stride": 4,
        "layer2_activation_function": "tanh",
        "layer2_filters": 15,
        "layer2_kernel_size": 6,
        "layer2_pool_size": 5,
        "layer2_stride": 2,
        "zeta_pos": 2.0,
        "zeta_mean": 0.1 ** 2,
        "zeta_var": 0.1 ** 2,
        "var_x": 0.005 ** 2,
        "var_y": 0.05 ** 2,
        "gp_var": 0.5 ** 2,
        "gp_len": 2.0,
        "random_seed": None,
        "in_dir": str(get_default_input_dir()),
        "output_root": str(get_default_output_root()),
        "output_name": None,
        "notes": "",
    }


def get_bootstrap_configuration(bootstrap_index: int, base_overrides: Dict[str, Any] | None = None) -> Dict[str, Any]:
    bootstrap_candidates: List[Dict[str, Any]] = [
        {
            "number_of_layers": 1,
            "layer1_activation_function": "sigmoid",
            "layer1_filters": 6,
            "layer1_kernel_size": 3,
            "layer1_pool_size": 2,
            "layer1_stride": 4,
            "learning_rate": 0.03,
            "use_learning_rate_decay": False,
            "batch_size": 14,
            "patience": 2,
            "min_delta": 3.162278,
            "recog_model": "zeros",
        },
        {
            "number_of_layers": 1,
            "layer1_activation_function": "ReLU",
            "layer1_filters": 8,
            "layer1_kernel_size": 6,
            "layer1_pool_size": 4,
            "layer1_stride": 2,
            "learning_rate": 0.019744,
            "use_learning_rate_decay": True,
            "batch_size": 8,
            "patience": 3,
            "min_delta": 0.153992,
            "recog_model": "zeros",
        },
        {
            "number_of_layers": 1,
            "layer1_activation_function": "tanh",
            "layer1_filters": 5,
            "layer1_kernel_size": 3,
            "layer1_pool_size": 2,
            "layer1_stride": 2,
            "learning_rate": 0.03,
            "use_learning_rate_decay": False,
            "batch_size": 5,
            "patience": 6,
            "min_delta": 0.056234,
            "recog_model": "conv",
        },
        {
            "number_of_layers": 1,
            "layer1_activation_function": "ReLU",
            "layer1_filters": 8,
            "layer1_kernel_size": 6,
            "layer1_pool_size": 4,
            "layer1_stride": 2,
            "learning_rate": 0.009487,
            "use_learning_rate_decay": False,
            "batch_size": 9,
            "patience": 4,
            "min_delta": 1.154781,
            "recog_model": "zeros",
        },
        {
            "number_of_layers": 1,
            "layer1_activation_function": "tanh",
            "layer1_filters": 8,
            "layer1_kernel_size": 6,
            "layer1_pool_size": 4,
            "layer1_stride": 2,
            "learning_rate": 0.094868,
            "use_learning_rate_decay": False,
            "batch_size": 11,
            "patience": 3,
            "min_delta": 0.542469,
            "recog_model": "zeros",
        },
    ]
    candidate = bootstrap_candidates[bootstrap_index % len(bootstrap_candidates)]
    overrides = dict(base_overrides or {})
    overrides.update(candidate)
    return build_experiment_config(overrides)


def get_bootstrap_recommendation(
    bootstrap_index: int,
    base_overrides: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    config = get_bootstrap_configuration(bootstrap_index, base_overrides=base_overrides)
    return {
        "configuration": {
            key: config.get(key)
            for key in FEATURE_COLUMNS
            if key in config
        },
        "is_novel_configuration": True,
        "predicted_average_rmse": None,
        "predicted_average_r2": None,
        "predicted_score_raw": None,
        "predicted_score": None,
        "overtraining_risk": "low",
        "overtraining_reason": (
            f"bootstrap campaign run {bootstrap_index + 1} of {MIN_HISTORY_FOR_SURROGATE} "
            "before enabling the full surrogate optimizer"
        ),
        "classification_preference": "bootstrap",
        "predicted_epoch": None,
        "predicted_training_time_seconds": None,
    }


def validate_runtime_environment() -> None:
    tf_version = getattr(tf, "__version__", "unknown")
    if not hasattr(tf, "placeholder"):
        raise RuntimeError(
            "This PRSSM automation requires the TensorFlow 1.x environment used by the "
            "manual training workflow. The current interpreter is '{}' with TensorFlow {}. "
            "Please run the automation with the prssm-fw environment, for example:\n"
            "  python Python_PRSSM/automation_runner.py --max-runs 1".format(
                sys.executable,
                tf_version,
            )
        )


def _coerce_int(value: Any, default: int) -> int:
    if value is None:
        return default
    return int(round(float(value)))


def _coerce_float(value: Any, default: float) -> float:
    if value is None:
        return default
    return float(value)


def _coerce_bool(value: Any, default: bool) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1", "yes", "y"}:
        return True
    if text in {"false", "0", "no", "n", ""}:
        return False
    return bool(value)


def _normalize_optional_numeric(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and value.strip().upper() in {"N/A", "NA", ""}:
        return None
    return value


def _normalize_optional_text(value: Any, default: str) -> str:
    if value is None:
        return default
    if isinstance(value, str) and value.strip().upper() in {"N/A", "NA", ""}:
        return default
    return str(value)


def _normalize_optional_int(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, str) and value.strip().upper() in {"N/A", "NA", ""}:
        return None
    return int(round(float(value)))


def _slug_text(value: Any) -> str:
    if value is None:
        return "NA"
    text = str(value).strip()
    if not text:
        return "NA"
    replacements = {
        "Exponential decay": "ED",
        "sigmoid": "sig",
        "ReLU": "relu",
        "tanh": "tanh",
        "N/A": "NA",
        " ": "",
        ".": "p",
        "-": "m",
        "/": "_",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return "".join(ch for ch in text if ch.isalnum() or ch == "_") or "NA"


def _classification_code(value: Any) -> str:
    text = str(value or "").strip().lower()
    if not text:
        return "PEND"
    if text == "undertrained":
        return "UT"
    if text == "overtrained" or text == "overtraining":
        return "OT"
    if text == "well-balanced" or text == "well balanced":
        return "WB"
    if text == "inconclusive":
        return "INC"
    return "RUN"


def build_experiment_config(overrides: Dict[str, Any] | None = None) -> Dict[str, Any]:
    config = get_default_experiment_config()
    if overrides:
        config.update(overrides)

    config["seq_len"] = _coerce_int(config.get("seq_len"), 470)
    config["seq_stride"] = _coerce_int(config.get("seq_stride"), config["seq_len"])
    config["epochs"] = _coerce_int(config.get("epochs"), 500)
    config["num_gpus"] = _coerce_int(config.get("num_gpus"), 1)
    config["batch_size"] = _coerce_int(config.get("batch_size"), 14)
    config["shuffle"] = _coerce_int(config.get("shuffle"), 10000)
    config["dim_x"] = _coerce_int(config.get("dim_x"), 12)
    config["ind_pnt_num"] = _coerce_int(config.get("ind_pnt_num"), 100)
    config["samples"] = _coerce_int(config.get("samples"), 30)
    config["patience"] = _coerce_int(config.get("patience"), 2)
    config["number_of_layers"] = max(1, _coerce_int(config.get("number_of_layers"), 1))
    config["recog_len"] = _coerce_int(config.get("recog_len"), 60)
    config["validation_fraction"] = _coerce_float(config.get("validation_fraction"), 0.1)
    config["random_seed"] = _normalize_optional_int(config.get("random_seed"))
    config["layer1_activation_function"] = _normalize_optional_text(
        config.get("layer1_activation_function"),
        "sigmoid",
    )
    config["layer2_activation_function"] = _normalize_optional_text(
        config.get("layer2_activation_function"),
        "tanh",
    )
    config["recog_model"] = _normalize_optional_text(config.get("recog_model"), "zeros")
    config["early_stopping_split"] = _normalize_optional_text(
        config.get("early_stopping_split"),
        "validation",
    ).lower()

    learning_rate_value = config.get("learning_rate")
    if isinstance(learning_rate_value, str) and learning_rate_value.strip().lower() in {
        "exponential decay",
        "exp decay",
        "exponential",
    }:
        config["learning_rate"] = 0.03
        config["use_learning_rate_decay"] = True

    for key, default in [
        ("learning_rate", 0.03),
        ("learning_rate_decay_steps", 1000.0),
        ("learning_rate_decay_rate", 0.96),
        ("min_delta", 3.162278),
        ("lik_seq_length_factor", 0.8),
        ("zeta_pos", 2.0),
        ("zeta_mean", 0.1 ** 2),
        ("zeta_var", 0.1 ** 2),
        ("var_x", 0.005 ** 2),
        ("var_y", 0.05 ** 2),
        ("gp_var", 0.5 ** 2),
        ("gp_len", 2.0),
    ]:
        config[key] = _coerce_float(config.get(key), default)

    config["use_learning_rate_decay"] = _coerce_bool(
        config.get("use_learning_rate_decay"),
        False,
    )
    config["learning_rate_decay_staircase"] = _coerce_bool(
        config.get("learning_rate_decay_staircase"),
        True,
    )
    config["test_data"] = _coerce_bool(config.get("test_data"), True)
    config["early_stopping"] = _coerce_bool(config.get("early_stopping"), True)

    for key, default in [
        ("layer1_filters", 6),
        ("layer1_kernel_size", 3),
        ("layer1_pool_size", 2),
        ("layer1_stride", 4),
        ("layer2_filters", 15),
        ("layer2_kernel_size", 6),
        ("layer2_pool_size", 5),
        ("layer2_stride", 2),
    ]:
        value = _normalize_optional_numeric(config.get(key))
        config[key] = None if value is None else _coerce_int(value, default)

    layer_keys = [
        "number_of_layers",
        "layer1_activation_function",
        "layer1_filters",
        "layer1_kernel_size",
        "layer1_pool_size",
        "layer1_stride",
        "layer2_activation_function",
        "layer2_filters",
        "layer2_kernel_size",
        "layer2_pool_size",
        "layer2_stride",
    ]
    if any(key in (overrides or {}) for key in layer_keys) and "recog_model" not in (overrides or {}):
        config["recog_model"] = "conv"

    if config["number_of_layers"] <= 1:
        # For single-layer runs, layer-2 settings are intentionally unused.
        # Keep them empty in saved configs/status so they cannot be mistaken as active.
        config["layer2_activation_function"] = None
        config["layer2_filters"] = None
        config["layer2_kernel_size"] = None
        config["layer2_pool_size"] = None
        config["layer2_stride"] = None

    config["gpus"] = [f"/device:GPU:{idx}" for idx in range(config["num_gpus"])]
    config["predict_len"] = int(config["seq_len"] * config["lik_seq_length_factor"])
    config["notes"] = str(config.get("notes", "")).strip()
    config["in_dir"] = str(Path(config["in_dir"]))
    config["output_root"] = str(Path(config["output_root"]))
    return config


def _normalize_single_layer_payload_fields(payload: Dict[str, Any]) -> Dict[str, Any]:
    normalized = dict(payload)
    config = normalized.get("configuration")
    if isinstance(config, dict):
        config = dict(config)
        if int(config.get("number_of_layers", 1) or 1) <= 1:
            config["layer2_activation_function"] = None
            config["layer2_filters"] = None
            config["layer2_kernel_size"] = None
            config["layer2_pool_size"] = None
            config["layer2_stride"] = None
        normalized["configuration"] = config

    recommendation = normalized.get("recommendation")
    if isinstance(recommendation, dict):
        recommendation = dict(recommendation)
        recommendation_config = recommendation.get("configuration")
        if isinstance(recommendation_config, dict):
            recommendation_config = dict(recommendation_config)
            if int(recommendation_config.get("number_of_layers", 1) or 1) <= 1:
                recommendation_config["layer2_activation_function"] = "N/A"
                recommendation_config["layer2_filters"] = "N/A"
                recommendation_config["layer2_kernel_size"] = "N/A"
                recommendation_config["layer2_pool_size"] = "N/A"
                recommendation_config["layer2_stride"] = "N/A"
            recommendation["configuration"] = recommendation_config
        normalized["recommendation"] = recommendation

    return normalized


def build_run_name(config: Dict[str, Any], prefix: str = "OF") -> str:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{prefix}_{timestamp}"


def build_descriptive_run_name(
    config: Dict[str, Any],
    training_test_number: int,
    epoch: Any = None,
    classification: Any = None,
) -> str:
    parts = [
        f"{int(training_test_number):02d}",
        f"L{_slug_text(config.get('number_of_layers'))}",
        _slug_text(config.get("layer1_activation_function")),
    ]
    if int(config.get("number_of_layers", 1) or 1) > 1:
        parts.append(f"L2{_slug_text(config.get('layer2_activation_function'))}")
    parts.extend(
        [
            f"lr{_slug_text(config.get('learning_rate'))}",
            f"b{_slug_text(config.get('batch_size'))}",
            f"p{_slug_text(config.get('patience'))}",
            f"d{_slug_text(config.get('min_delta'))}",
            f"e{_slug_text(epoch if epoch is not None else 'PRE')}",
            _classification_code(classification),
        ]
    )
    return f"{parts[0]} - {'_'.join(parts[1:])}"


def build_temporary_run_name(training_test_number: int) -> str:
    return f"R{int(training_test_number):02d}"


def resolve_output_dir(config: Dict[str, Any], output_name: str | None = None) -> Path:
    output_root = Path(config["output_root"])
    final_name = output_name or config.get("output_name") or build_run_name(config)
    out_dir = output_root / final_name
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir


def build_model_config(config: Dict[str, Any], out_dir: Path) -> Dict[str, Any]:
    return {
        "batch_size": config["batch_size"],
        "shuffle": config["shuffle"],
        "lik_seq_length_factor": config["lik_seq_length_factor"],
        "dim_x": config["dim_x"],
        "ind_pnt_num": config["ind_pnt_num"],
        "samples": config["samples"],
        "learning_rate": config["learning_rate"],
        "use_learning_rate_decay": config["use_learning_rate_decay"],
        "learning_rate_decay_steps": config["learning_rate_decay_steps"],
        "learning_rate_decay_rate": config["learning_rate_decay_rate"],
        "learning_rate_decay_staircase": config["learning_rate_decay_staircase"],
        "recog_len": config["recog_len"],
        "recog_model": config["recog_model"],
        "number_of_layers": config["number_of_layers"],
        "layer1_activation_function": config["layer1_activation_function"],
        "layer1_filters": config["layer1_filters"],
        "layer1_kernel_size": config["layer1_kernel_size"],
        "layer1_pool_size": config["layer1_pool_size"],
        "layer1_stride": config["layer1_stride"],
        "layer2_activation_function": config["layer2_activation_function"],
        "layer2_filters": config["layer2_filters"],
        "layer2_kernel_size": config["layer2_kernel_size"],
        "layer2_pool_size": config["layer2_pool_size"],
        "layer2_stride": config["layer2_stride"],
        "patience": config["patience"],
        "min_delta": config["min_delta"],
        "early_stopping_split": config["early_stopping_split"],
        "validation_fraction": config["validation_fraction"],
        "zeta_pos": config["zeta_pos"],
        "zeta_mean": config["zeta_mean"],
        "zeta_var": config["zeta_var"],
        "var_x": config["var_x"],
        "var_y": config["var_y"],
        "gp_var": config["gp_var"],
        "gp_len": config["gp_len"],
        "random_seed": config["random_seed"],
        "gpus": config["gpus"],
        "epochs": config["epochs"],
        "in_dir": config["in_dir"],
        "out_dir": str(out_dir) + "/",
        "notes": config.get("notes", ""),
    }


def apply_random_seed(random_seed: int | None) -> None:
    if random_seed is None:
        return

    try:
        import random

        random.seed(random_seed)
    except Exception:
        pass

    try:
        import numpy as np

        np.random.seed(random_seed)
    except Exception:
        pass

    try:
        tf.set_random_seed(random_seed)
    except AttributeError:
        pass
    try:
        tf.compat.v1.set_random_seed(random_seed)
    except AttributeError:
        pass


def save_training_config(model_config: Dict[str, Any], out_dir: Path) -> Path:
    config_path = out_dir / "training_config.json"
    with config_path.open("w", encoding="utf-8") as fout:
        json.dump(model_config, fout, indent=2)
    return config_path


def load_training_config(config_path: Path) -> Dict[str, Any]:
    with config_path.open("r", encoding="utf-8") as fin:
        return json.load(fin)


def summarize_metrics(metrics_file: Path) -> Dict[str, Any]:
    if not metrics_file.exists():
        return {}

    metrics_df = pd.read_csv(metrics_file)
    metrics_df = metrics_df.copy()
    metrics_df["split"] = metrics_df["split"].astype(str).str.lower()

    summary: Dict[str, Any] = {}
    test_rows = metrics_df[
        (metrics_df["split"] == "test") & (metrics_df["parity_plot_version"] == "perfect_line")
    ].copy()
    if not test_rows.empty:
        summary["average_test_rmse"] = float(test_rows["rmse"].mean())
        summary["average_test_r2"] = float(test_rows["r2"].mean())
        for output_name in OUTPUT_COLUMNS:
            variable_name = output_name.replace("_", "_")
            source_name = output_name.replace("_", "_")
            target_row = test_rows[test_rows["variable"].astype(str).str.upper() == output_name]
            if target_row.empty:
                target_row = test_rows[test_rows["variable"].astype(str).str.replace("_", "", regex=False).str.upper() == output_name.replace("_", "")]
            if target_row.empty:
                continue
            row = target_row.iloc[0]
            summary[f"{source_name}_RMSE"] = float(row["rmse"])
            summary[f"{source_name}_R2"] = float(row["r2"])
    return summary


def summarize_training_log(log_path: Path) -> Dict[str, Any]:
    if not log_path.exists():
        return {}

    history_df, log_summary = parse_training_log(log_path)
    result: Dict[str, Any] = {}
    if "total_training_time_sec" in log_summary:
        result["training_time_seconds"] = float(log_summary["total_training_time_sec"])
        result["training_time_hours"] = float(log_summary["total_training_time_sec"]) / 3600.0
    if "best_test_epoch" in log_summary:
        result["epoch"] = int(log_summary["best_test_epoch"])
    if "best_test_loss" in log_summary:
        result["best_test_loss"] = float(log_summary["best_test_loss"])
    if "training_classification" in log_summary:
        result["training_classification"] = str(log_summary["training_classification"]).strip().lower()
    if "training_classification_reason" in log_summary:
        result["training_classification_reason"] = str(log_summary["training_classification_reason"]).strip()
    if not history_df.empty:
        result["epochs_logged"] = int(len(history_df))
    return result


def build_experiment_row(config: Dict[str, Any], run_result: Dict[str, Any]) -> Dict[str, Any]:
    row: Dict[str, Any] = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "run_name": run_result.get("run_name", ""),
        "output_dir": run_result.get("out_dir", ""),
        "notes": config.get("notes", ""),
    }

    for column in FEATURE_COLUMNS:
        row[column] = config.get(column)

    for output_name in OUTPUT_COLUMNS:
        row[f"{output_name}_RMSE"] = run_result.get(f"{output_name}_RMSE")
        row[f"{output_name}_R2"] = run_result.get(f"{output_name}_R2")

    row["average_test_rmse"] = run_result.get("average_test_rmse")
    row["average_test_r2"] = run_result.get("average_test_r2")
    row["epoch"] = run_result.get("epoch")
    row["total_training_time_seconds"] = run_result.get("training_time_seconds")
    row["total_training_time_hours"] = run_result.get("training_time_hours")
    row["training_classification"] = run_result.get("training_classification", "")
    row["training_classification_reason"] = run_result.get("training_classification_reason", "")
    row["best_test_loss"] = run_result.get("best_test_loss")
    row["best_test_epoch"] = run_result.get("epoch")
    return row


def describe_split_and_seed_policy(config: Dict[str, Any]) -> str:
    split_text = (
        "split_policy=fixed trajectories in train/validation/test for each run"
    )
    random_seed = config.get("random_seed")
    if random_seed is None:
        seed_text = "seed_policy=randomized each run"
    else:
        seed_text = f"seed_policy=fixed seed {int(random_seed)} each run"
    return f"{split_text}; {seed_text}"


def get_summary_internal_headers(worksheet) -> List[str]:
    top = pd.Series([worksheet.cell(row=1, column=i + 1).value for i in range(worksheet.max_column)])
    bottom = pd.Series([worksheet.cell(row=2, column=i + 1).value for i in range(worksheet.max_column)])
    return build_column_names(top, bottom)


def get_training_test_number_column_index(worksheet) -> int:
    internal_headers = get_summary_internal_headers(worksheet)
    if TRAINING_TEST_NUMBER_INTERNAL_COLUMN not in internal_headers:
        raise KeyError(
            "The Summary sheet does not contain the expected 'Training Test Number' column."
        )
    return internal_headers.index(TRAINING_TEST_NUMBER_INTERNAL_COLUMN) + 1


def get_next_training_test_number(excel_path: Path, sheet_name: str) -> int:
    existing_values: List[int] = []
    if load_workbook is None or not excel_path.exists():
        pass
    else:
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            if sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]
                training_col_idx = get_training_test_number_column_index(worksheet)
                for row_idx in range(3, worksheet.max_row + 1):
                    cell_value = worksheet.cell(row=row_idx, column=training_col_idx).value
                    if cell_value in {None, ""}:
                        continue
                    try:
                        existing_values.append(int(float(cell_value)))
                    except (TypeError, ValueError):
                        continue
        finally:
            workbook.close()

    output_root = get_default_output_root()
    if output_root.exists():
        for child in output_root.iterdir():
            if not child.is_dir():
                continue
            prefix = child.name.split(" - ", 1)[0].strip()
            if prefix.upper().startswith("R") and prefix[1:].isdigit():
                existing_values.append(int(prefix[1:]))
                continue
            try:
                existing_values.append(int(prefix))
            except (TypeError, ValueError):
                continue
    return (max(existing_values) + 1) if existing_values else 1


def count_logged_runs(excel_path: Path, sheet_name: str) -> int:
    if load_workbook is None or not excel_path.exists():
        return 0
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return 0
        worksheet = workbook[sheet_name]
        training_col_idx = get_training_test_number_column_index(worksheet)
        count = 0
        for row_idx in range(3, worksheet.max_row + 1):
            cell_value = worksheet.cell(row=row_idx, column=training_col_idx).value
            if cell_value in {None, ""}:
                continue
            try:
                float(cell_value)
                count += 1
            except (TypeError, ValueError):
                continue
        return count
    finally:
        workbook.close()


def append_row_to_summary_sheet(workbook, worksheet, row: Dict[str, Any]) -> None:
    internal_headers = get_summary_internal_headers(worksheet)
    training_col_idx = get_training_test_number_column_index(worksheet)
    next_row_idx = worksheet.max_row + 1
    for col_idx, internal_header in enumerate(internal_headers, start=1):
        row_key = SUMMARY_ROW_KEY_BY_INTERNAL_HEADER.get(internal_header)
        worksheet.cell(row=next_row_idx, column=col_idx).value = row.get(row_key) if row_key else None
    worksheet.cell(row=next_row_idx, column=training_col_idx).value = row.get("training_test_number")


def append_row_to_excel(excel_path: Path, row: Dict[str, Any], sheet_name: str = DEFAULT_AUTOMATION_SHEET) -> None:
    if Workbook is None or load_workbook is None:
        raise ImportError("Writing to Excel requires openpyxl to be installed.")

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    if excel_path.exists():
        workbook = load_workbook(excel_path)
    else:
        workbook = Workbook()

    if sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.create_sheet(title=sheet_name)
        if worksheet.max_row == 1 and worksheet.max_column == 1 and worksheet["A1"].value is None:
            worksheet.delete_rows(1, 1)

    if sheet_name == "Summary" and worksheet.max_row >= 2:
        append_row_to_summary_sheet(workbook, worksheet, row)
        workbook.save(excel_path)
        return

    headers = list(row.keys())
    if worksheet.max_row == 0:
        worksheet.append(headers)
    elif worksheet.max_row == 1 and worksheet.max_column == 1 and worksheet["A1"].value is None:
        worksheet.append(headers)
    elif [worksheet.cell(row=1, column=i + 1).value for i in range(len(headers))] != headers:
        existing_headers = [worksheet.cell(row=1, column=i + 1).value for i in range(worksheet.max_column)]
        if existing_headers != headers:
            worksheet = workbook.create_sheet(title=f"{sheet_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
            worksheet.append(headers)

    worksheet.append([row.get(header) for header in headers])
    workbook.save(excel_path)


def resolve_workbook_sheet_name(
    excel_path: Path,
    requested_sheet_name: str | None = None,
    fallback_sheet_name: str = DEFAULT_AUTOMATION_SHEET,
) -> str:
    if requested_sheet_name:
        return requested_sheet_name
    if excel_path.exists() and load_workbook is not None:
        workbook = load_workbook(excel_path, read_only=True)
        try:
            if "Summary" in workbook.sheetnames:
                return "Summary"
            if workbook.sheetnames:
                return workbook.sheetnames[0]
        finally:
            workbook.close()
    return fallback_sheet_name


def save_experiment_summary(out_dir: Path, summary: Dict[str, Any]) -> Path:
    summary_path = out_dir / "automation_summary.json"
    with summary_path.open("w", encoding="utf-8") as fout:
        json.dump(summary, fout, indent=2)
    return summary_path


def save_current_run_status(base_path: Path, payload: Dict[str, Any]) -> Path:
    base_path.mkdir(parents=True, exist_ok=True)
    status_path = base_path / CURRENT_RUN_STATUS_FILE
    payload = _normalize_single_layer_payload_fields(payload)
    with status_path.open("w", encoding="utf-8") as fout:
        json.dump(payload, fout, indent=2)
    return status_path


def clear_current_run_status(base_path: Path) -> None:
    status_path = base_path / CURRENT_RUN_STATUS_FILE
    if status_path.exists():
        status_path.unlink()


def start_run_heartbeat(
    base_path: Path,
    payload: Dict[str, Any],
    interval_seconds: int = HEARTBEAT_INTERVAL_SECONDS,
) -> tuple[threading.Event, threading.Thread]:
    stop_event = threading.Event()

    def heartbeat_loop() -> None:
        while not stop_event.is_set():
            heartbeat_payload = dict(payload)
            heartbeat_payload["status"] = "running"
            heartbeat_payload["last_heartbeat_timestamp"] = datetime.now().isoformat(timespec="seconds")
            save_current_run_status(base_path, heartbeat_payload)
            stop_event.wait(interval_seconds)

    thread = threading.Thread(target=heartbeat_loop, name="automation-heartbeat", daemon=True)
    thread.start()
    return stop_event, thread


def stop_request_file_path(base_path: Path) -> Path:
    return base_path / STOP_REQUEST_FILE


def is_stop_requested(
    base_path: Path,
    stop_requested: Callable[[], bool] | None = None,
) -> bool:
    callback_requested = bool(stop_requested and stop_requested())
    file_requested = stop_request_file_path(base_path).exists()
    return callback_requested or file_requested


def format_parameter_summary(config: Dict[str, Any]) -> str:
    lines = [
        f"output_name: {config.get('output_name') or build_run_name(config)}",
        f"recog_model: {config['recog_model']}",
        f"number_of_layers: {config['number_of_layers']}",
        f"layer1_activation_function: {config['layer1_activation_function']}",
        f"layer1_filters: {config['layer1_filters']}",
        f"layer1_kernel_size: {config['layer1_kernel_size']}",
        f"layer1_pool_size: {config['layer1_pool_size']}",
        f"layer1_stride: {config['layer1_stride']}",
        f"layer2_activation_function: {config['layer2_activation_function']}",
        f"layer2_filters: {config['layer2_filters']}",
        f"layer2_kernel_size: {config['layer2_kernel_size']}",
        f"layer2_pool_size: {config['layer2_pool_size']}",
        f"layer2_stride: {config['layer2_stride']}",
        f"learning_rate: {config['learning_rate']}",
        f"random_seed: {config.get('random_seed')}",
        f"use_learning_rate_decay: {config['use_learning_rate_decay']}",
        f"batch_size: {config['batch_size']}",
        f"patience: {config['patience']}",
        f"min_delta: {config['min_delta']}",
        f"early_stopping_split: {config['early_stopping_split']}",
        f"validation_fraction: {config['validation_fraction']}",
        f"epochs: {config['epochs']}",
        f"notes: {config.get('notes', '')}",
    ]
    return "\n".join(lines)


def refresh_metadata_paths(out_dir: Path, result: Dict[str, Any]) -> None:
    training_config_path = out_dir / "training_config.json"
    if training_config_path.exists():
        training_config = load_training_config(training_config_path)
        training_config["out_dir"] = str(out_dir) + "/"
        with training_config_path.open("w", encoding="utf-8") as fout:
            json.dump(training_config, fout, indent=2)

    result["run_name"] = out_dir.name
    result["out_dir"] = str(out_dir)
    if isinstance(result.get("config"), dict):
        result["config"]["output_name"] = out_dir.name
        result["config"]["output_root"] = str(out_dir.parent)
        result["config"]["out_dir"] = str(out_dir)
    if isinstance(result.get("postprocess"), dict):
        result["postprocess"]["out_dir"] = str(out_dir)
        metrics_path = out_dir / "metrics_summary.csv"
        if metrics_path.exists():
            result["postprocess"]["metrics_file"] = str(metrics_path)


def rename_run_directory(
    out_dir: Path,
    config: Dict[str, Any],
    training_test_number: int,
    run_result: Dict[str, Any],
) -> Path:
    final_name = build_descriptive_run_name(
        config=config,
        training_test_number=training_test_number,
        epoch=run_result.get("epoch"),
        classification=run_result.get("training_classification"),
    )
    final_dir = out_dir.parent / final_name
    if final_dir == out_dir:
        refresh_metadata_paths(final_dir, run_result)
        save_experiment_summary(final_dir, run_result)
        return final_dir
    if final_dir.exists():
        raise FileExistsError(f"Cannot rename run directory because the target already exists: {final_dir}")
    out_dir.rename(final_dir)
    refresh_metadata_paths(final_dir, run_result)
    save_experiment_summary(final_dir, run_result)
    return final_dir


def is_risk_allowed(recommendation: Dict[str, Any], allowed_risks: Sequence[str] | None) -> bool:
    if not allowed_risks:
        return True
    allowed = {str(risk).strip().lower() for risk in allowed_risks}
    current = str(recommendation.get("overtraining_risk", "")).strip().lower()
    return current in allowed


def run_experiment(
    overrides: Dict[str, Any] | None = None,
    output_name: str | None = None,
    retrain: bool = False,
    train: bool = True,
    run_postprocess: bool = True,
) -> Dict[str, Any]:
    validate_runtime_environment()
    config = build_experiment_config(overrides)
    out_dir = resolve_output_dir(config, output_name=output_name)
    config["output_name"] = out_dir.name
    model_dir = str(out_dir) + "/"
    model_config = build_model_config(config, out_dir)
    save_training_config(model_config, out_dir)
    print("\nStarting experiment with parameters:")
    print(format_parameter_summary(config))

    apply_random_seed(config.get("random_seed"))

    ds = FlappingWingAerodynamics(
        config["seq_len"],
        config["seq_stride"],
        config["in_dir"],
        validation_fraction=float(config["validation_fraction"]),
    )
    outputs = Outputs(model_dir)
    outputs.set_ds(ds)

    model = PRSSM(ds.dim_u, ds.dim_y, model_config)
    outputs.set_model(model, config["predict_len"], model_config["dim_x"])

    if train:
        trainer = Trainer(model, model_dir)
        trainer.train(
            ds,
            config["epochs"],
            retrain=retrain,
            test_data=bool(config["test_data"]),
            early_stopping=bool(config["early_stopping"]),
            patience=int(config["patience"]),
            min_delta=float(config["min_delta"]),
            evaluation_split=str(config["early_stopping_split"]),
        )
        outputs.set_trainer(trainer)

    outputs.create_all()

    postprocess_result: Dict[str, Any] = {}
    if run_postprocess:
        postprocess_result = run_for_output_dir(out_dir, show_plots=False)

    result: Dict[str, Any] = {
        "run_name": out_dir.name,
        "out_dir": str(out_dir),
        "config": config,
        "postprocess": postprocess_result,
    }
    result.update(summarize_metrics(out_dir / "metrics_summary.csv"))
    result.update(summarize_training_log(out_dir / "training_log.txt"))
    result.update({
        key: postprocess_result.get(key)
        for key in ["training_classification", "training_classification_reason", "best_test_loss", "best_test_epoch"]
        if key in postprocess_result
    })
    save_experiment_summary(out_dir, result)
    return result


def recommendation_to_overrides(recommendation: Dict[str, Any], base_overrides: Dict[str, Any] | None = None) -> Dict[str, Any]:
    overrides = dict(base_overrides or {})
    overrides.update(recommendation.get("configuration", {}))
    # Keep epochs as a generous training budget so early stopping, not the suggested
    # epoch count, determines when training ends.
    overrides.pop("epochs", None)
    if any(key in overrides for key in FEATURE_COLUMNS):
        overrides.setdefault("recog_model", "conv")
    return overrides


def select_strategy_for_run(
    run_index: int,
    default_strategy: str,
    exploratory_runs: int = 0,
    exploratory_strategy: str = "best_accuracy_recommendation",
) -> str:
    if run_index < max(0, exploratory_runs):
        return exploratory_strategy
    return default_strategy


def build_run_notes(
    run_index: int,
    max_runs: int,
    effective_strategy: str,
    recommendation: Dict[str, Any],
    config: Dict[str, Any],
    exploratory_runs: int = 0,
) -> str:
    phase = "exploration" if run_index < max(0, exploratory_runs) else "exploitation"
    predicted_rmse = recommendation.get("predicted_average_rmse")
    predicted_r2 = recommendation.get("predicted_average_r2")
    risk = recommendation.get("overtraining_risk", "")
    novelty = recommendation.get("is_novel_configuration")

    if phase == "exploration":
        reason = (
            "Early exploratory run to probe a high-upside region and expand coverage before "
            "switching to balanced follow-up recommendations."
        )
    else:
        reason = (
            "Balanced follow-up run to refine around promising low-risk settings after the initial "
            "exploratory phase."
        )

    details = [
        f"phase={phase}",
        f"run={run_index + 1}/{max_runs}",
        f"strategy={effective_strategy}",
        f"novel={novelty}",
        f"risk={risk}",
    ]
    if predicted_rmse is not None:
        details.append(f"pred_rmse={float(predicted_rmse):.4f}")
    if predicted_r2 is not None:
        details.append(f"pred_r2={float(predicted_r2):.4f}")
    details.append(describe_split_and_seed_policy(config))
    return reason + " " + "; ".join(details)


def unique_strategy_order(strategies: Sequence[str]) -> List[str]:
    ordered: List[str] = []
    for strategy in strategies:
        if strategy and strategy not in ordered:
            ordered.append(strategy)
    return ordered


def propose_next_run(
    excel_path: Path = DEFAULT_EXCEL_PATH,
    sheet_name: str | None = None,
    strategy: str = "best_balanced_accuracy_recommendation",
    exclude_overtraining: bool = False,
    base_overrides: Dict[str, Any] | None = None,
) -> Dict[str, Any]:
    print(f"experiment_runner: resolving workbook sheet for {excel_path}", flush=True)
    resolved_sheet_name = resolve_workbook_sheet_name(excel_path, sheet_name)
    print(f"experiment_runner: resolved sheet '{resolved_sheet_name}'", flush=True)
    logged_runs = count_logged_runs(excel_path, resolved_sheet_name)
    print(f"experiment_runner: logged runs in workbook = {logged_runs}", flush=True)
    if logged_runs < MIN_HISTORY_FOR_SURROGATE:
        recommendation = get_bootstrap_recommendation(
            bootstrap_index=logged_runs,
            base_overrides=base_overrides,
        )
        return {
            "optimizer_results": {"bootstrap_recommendation": recommendation},
            "resolved_sheet_name": resolved_sheet_name,
            "selected_strategy": strategy,
            "recommendation": recommendation,
            "overrides": recommendation_to_overrides(recommendation, base_overrides=base_overrides),
        }
    results = optimize_hyperparameters(
        excel_path=excel_path,
        sheet_name=resolved_sheet_name,
        exclude_overtraining=exclude_overtraining,
    )
    if strategy not in results:
        raise KeyError(f"Unknown recommendation strategy: {strategy}")
    recommendation = results[strategy]
    return {
        "optimizer_results": results,
        "resolved_sheet_name": resolved_sheet_name,
        "selected_strategy": strategy,
        "recommendation": recommendation,
        "overrides": recommendation_to_overrides(recommendation, base_overrides=base_overrides),
    }


def propose_next_allowed_run(
    excel_path: Path,
    sheet_name: str | None,
    preferred_strategies: Sequence[str],
    exclude_overtraining: bool,
    base_overrides: Dict[str, Any] | None,
    allowed_risks: Sequence[str] | None,
) -> Dict[str, Any]:
    last_proposal: Dict[str, Any] | None = None
    for candidate_strategy in unique_strategy_order(preferred_strategies):
        proposal = propose_next_run(
            excel_path=excel_path,
            sheet_name=sheet_name,
            strategy=candidate_strategy,
            exclude_overtraining=exclude_overtraining,
            base_overrides=base_overrides,
        )
        proposal["selected_strategy"] = candidate_strategy
        last_proposal = proposal
        if is_risk_allowed(proposal["recommendation"], allowed_risks):
            return proposal
    if last_proposal is None:
        raise RuntimeError("No recommendation strategies were available to evaluate.")
    return last_proposal


def run_automation_cycle(
    max_runs: int = 1,
    excel_path: Path = DEFAULT_EXCEL_PATH,
    sheet_name: str | None = None,
    strategy: str = "best_balanced_accuracy_recommendation",
    exploratory_runs: int = 0,
    exploratory_strategy: str = "best_accuracy_recommendation",
    exclude_overtraining: bool = False,
    base_overrides: Dict[str, Any] | None = None,
    stop_requested: Callable[[], bool] | None = None,
    allowed_risks: Sequence[str] | None = ("low",),
) -> List[Dict[str, Any]]:
    validate_runtime_environment()
    results: List[Dict[str, Any]] = []
    status_base_path = get_automation_status_root()
    clear_current_run_status(status_base_path)
    print(f"experiment_runner: resolving logging sheet for {excel_path}", flush=True)
    logging_sheet_name = resolve_workbook_sheet_name(excel_path, sheet_name)
    print(f"experiment_runner: using logging sheet '{logging_sheet_name}'", flush=True)

    for run_index in range(max_runs):
        print(f"experiment_runner: cycle {run_index + 1} of {max_runs}", flush=True)
        if is_stop_requested(status_base_path, stop_requested):
            print("Stop requested before launching the next automation cycle. Exiting cleanly.")
            break
        effective_strategy = select_strategy_for_run(
            run_index=run_index,
            default_strategy=strategy,
            exploratory_runs=exploratory_runs,
            exploratory_strategy=exploratory_strategy,
        )
        proposal = propose_next_allowed_run(
            excel_path=excel_path,
            sheet_name=sheet_name,
            preferred_strategies=[
                effective_strategy,
                strategy,
                "best_conservative_recommendation",
                "best_balanced_accuracy_recommendation",
            ],
            exclude_overtraining=exclude_overtraining,
            base_overrides=base_overrides,
            allowed_risks=allowed_risks,
        )
        effective_strategy = proposal["selected_strategy"]
        recommendation = proposal["recommendation"]
        predicted_hours = recommendation.get("predicted_training_time_hours")
        print("\nNext recommended run:")
        print("  strategy: {}".format(effective_strategy))
        print("  predicted_score: {}".format(recommendation.get("predicted_score")))
        print("  predicted_average_r2: {}".format(recommendation.get("predicted_average_r2")))
        print("  predicted_average_rmse: {}".format(recommendation.get("predicted_average_rmse")))
        print("  overtraining_risk: {}".format(recommendation.get("overtraining_risk")))
        if predicted_hours is not None:
            print("  predicted_training_time_hours: {:.3f}".format(float(predicted_hours)))
        if not is_risk_allowed(recommendation, allowed_risks):
            print(
                "Skipping run because overtraining risk '{}' is outside allowed risks {}.".format(
                    recommendation.get("overtraining_risk"),
                    list(allowed_risks),
                )
            )
            break
        overrides = recommendation_to_overrides(
            recommendation,
            base_overrides=base_overrides,
        )
        training_test_number = get_next_training_test_number(excel_path, logging_sheet_name)
        overrides["notes"] = build_run_notes(
            run_index=run_index,
            max_runs=max_runs,
            effective_strategy=effective_strategy,
            recommendation=recommendation,
            config=build_experiment_config(overrides),
            exploratory_runs=exploratory_runs,
        )
        config = build_experiment_config(overrides)
        initial_output_name = build_temporary_run_name(training_test_number)
        config["output_name"] = initial_output_name
        current_run_payload = {
            "automation_cycle_index": run_index + 1,
            "training_test_number": training_test_number,
            "strategy": effective_strategy,
            "excel_path": str(excel_path),
            "sheet_name": logging_sheet_name,
            "status": "starting",
            "configuration": config,
            "recommendation": proposal["recommendation"],
        }
        save_current_run_status(status_base_path, current_run_payload)
        heartbeat_stop_event, heartbeat_thread = start_run_heartbeat(
            status_base_path,
            current_run_payload,
        )
        try:
            run_result = run_experiment(
                overrides=overrides,
                output_name=initial_output_name,
                train=True,
                run_postprocess=True,
            )
            final_out_dir = rename_run_directory(
                Path(run_result["out_dir"]),
                config,
                training_test_number,
                run_result,
            )
            run_result["run_name"] = final_out_dir.name
            run_result["out_dir"] = str(final_out_dir)
            logged_config = build_experiment_config(overrides)
            logged_config["output_name"] = final_out_dir.name
            row = build_experiment_row(logged_config, run_result)
            row["training_test_number"] = training_test_number
            append_row_to_excel(excel_path, row, sheet_name=logging_sheet_name)
            results.append(
                {
                    "proposal": proposal,
                    "run_result": run_result,
                    "logged_row": row,
                    "logging_sheet_name": logging_sheet_name,
                }
            )
            print("Saved completed run to workbook sheet '{}'.".format(logging_sheet_name))
            completion_payload = dict(current_run_payload)
            completion_payload["status"] = "completed"
            completion_payload["last_heartbeat_timestamp"] = datetime.now().isoformat(timespec="seconds")
            completion_payload["completed_run_name"] = run_result.get("run_name")
            completion_payload["completed_out_dir"] = run_result.get("out_dir")
            save_current_run_status(status_base_path, completion_payload)
        except Exception:
            failure_payload = dict(current_run_payload)
            failure_payload["status"] = "failed"
            failure_payload["last_heartbeat_timestamp"] = datetime.now().isoformat(timespec="seconds")
            save_current_run_status(status_base_path, failure_payload)
            raise
        finally:
            heartbeat_stop_event.set()
            heartbeat_thread.join(timeout=2.0)
            clear_current_run_status(status_base_path)
        if is_stop_requested(status_base_path, stop_requested):
            print("Stop requested. The completed run was saved, and no new run will be started.")
            break
    return results
